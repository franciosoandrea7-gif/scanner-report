import streamlit as st
import pandas as pd
import os
import random
import requests
import smtplib
from email.mime.multipart import MIMEMultipart
from email.mime.text import MIMEText
from email.mime.base import MIMEBase
from email import encoders
from PIL import Image
from datetime import datetime
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from streamlit_js_eval import get_geolocation

st.set_page_config(page_title="Nova Report Pro", page_icon="⚙️", layout="centered")
st.title("🛠️ Nova Report Pro")
st.write("Gestionale riparazioni Nova Servimpianti con validazione forte, foto multiple, GPS ed invio Email.")

EXCEL_FILE = "registro_riparazioni.xlsx"
LOGO_FILE = "logo.png"  

# --- CONFIGURAZIONE TEAM TECNICI E PIN SEGRETI ---
TECNICI = {
    "Andrea Francioso": "1974",
    "Daniele Gennari 1": "1990",
    "Lidia Distratis 2": "1977",
    "Nome Dipendente 3": "3415",
    "Nome Dipendente 4": "7712"
}

if "codice_sms" not in st.session_state:
    st.session_state["codice_sms"] = None
if "sms_validato" not in st.session_state:
    st.session_state["sms_validato"] = False
if "ultimo_pdf" not in st.session_state:
    st.session_state["ultimo_pdf"] = None
if "mostra_download" not in st.session_state:
    st.session_state["mostra_download"] = False

# --- RECUPERO DINAMICO LISTA CLIENTI DA EXCEL ---
lista_clienti_esistenti = []
if os.path.exists(EXCEL_FILE):
    try:
        wb = load_workbook(EXCEL_FILE, read_only=True)
        lista_clienti_esistenti = [sheet.replace("_", " ") for sheet in wb.sheetnames if sheet != "Sheet1"]
        wb.close()
    except Exception:
        lista_clienti_esistenti = []

lista_clienti_esistenti.sort()
opzioni_menu_clienti = ["➕ AGGIUNGI NUOVO CLIENTE"] + lista_clienti_esistenti

# --- 0. ARCHIVIO EXCEL ---
with st.expander("📚 Archivio Storico Lavori (Excel)"):
    if os.path.exists(EXCEL_FILE):
        try:
            wb = load_workbook(EXCEL_FILE, read_only=True)
            fogli = wb.sheetnames
            wb.close()
            if fogli:
                foglio_scelto = st.selectbox("Seleziona il registro del cliente da visualizzare:", fogli, key="view_select_client")
                st.dataframe(pd.read_excel(EXCEL_FILE, sheet_name=foglio_scelto))
            else:
                st.info("ℹ️ L'archivio Excel non contiene fogli validi.")
        except Exception as e:
            st.error(f"Errore nella lettura del registro: {e}")
    else:
        st.info("ℹ️ L'archivio Excel è vuoto.")

# --- 0. ARCHIVIO PDF CON RICERCA VELOCE ---
with st.expander("📂 Recupera Vecchi Report PDF Emessi"):
    lista_pdf = [f for f in os.listdir(".") if f.startswith("Report_") and f.endswith(".pdf")]
    if len(lista_pdf) > 0:
        lista_pdf.sort(reverse=True)
        cerca_pdf = st.text_input("🔍 Cerca PDF per nome cliente:", "").strip().lower()
        for nome_pdf in lista_pdf:
            if cerca_pdf and cerca_pdf not in nome_pdf.lower():
                continue
            col_n, col_b = st.columns(2)
            with col_n:
                st.write(f"📄 {nome_pdf.replace('Report_', '').replace('.pdf', '')}")
            with col_b:
                with open(nome_pdf, "rb") as f_pdf:
                    st.download_button("📥 Scarica", f_pdf, file_name=nome_pdf, key=f"st_{nome_pdf}")
    else:
        st.info("ℹ️ Nessun PDF in archivio.")

# --- SEZIONE SELEZIONE E FIRMA TECNICO ---
st.subheader("👨🔧 Responsabile Intervento")
tecnico_selezionato = st.selectbox("Seleziona il tuo nome dal personale Nova:", list(TECNICI.keys()))
pin_tecnico = st.text_input("Inserisci il tuo PIN Tecnico segreto per firmare:", type="password")

tecnico_autorizzato = False
if pin_tecnico:
    if TECNICI.get(tecnico_selezionato) == pin_tecnico:
        tecnico_autorizzato = True
        st.success(f"✍️ ID Tecnico Verificato: {tecnico_selezionato}")
    else:
        st.error("❌ PIN errato!")
        
# --- 1. MODULO DATI CLIENTE ---
st.subheader("👤 Dati Intervento")
data_corrente = st.date_input("Data Intervento", datetime.now())

# Funzione per forzare l'autocompilazione immediata nello schermo di Streamlit appena cambia il cliente
def aggiorna_dati_cliente():
    scelta = st.session_state["main_select_client"]
    if scelta != "➕ AGGIUNGI NUOVO CLIENTE":
        st.session_state["email_manual_input"] = ""
        st.session_state["cellulare_manual_input"] = ""
        if os.path.exists(EXCEL_FILE):
            try:
                nome_foglio = scelta.replace(" ", "_").replace("/", "_").replace("\\", "_").replace("?", "_").replace("*", "_")[:30]
                df_storico = pd.read_excel(EXCEL_FILE, sheet_name=nome_foglio)
                if not df_storico.empty:
                    if "Email Cliente" in df_storico.columns and pd.notna(df_storico["Email Cliente"].iloc[-1]):
                        st.session_state["email_manual_input"] = str(df_storico["Email Cliente"].iloc[-1]).strip()
                    if "Cellulare Cliente" in df_storico.columns and pd.notna(df_storico["Cellulare Cliente"].iloc[-1]):
                        st.session_state["cellulare_manual_input"] = str(df_storico["Cellulare Cliente"].iloc[-1]).strip()
            except Exception:
                pass
    else:
        st.session_state["email_manual_input"] = ""
        st.session_state["cellulare_manual_input"] = ""

# Menu a tendina del cliente con la funzione on_change attivata
cliente_selezionato_menu = st.selectbox("Seleziona Cliente *", opzioni_menu_clienti, key="main_select_client", on_change=aggiorna_dati_cliente)

if cliente_selezionato_menu == "➕ AGGIUNGI NUOVO CLIENTE":
    nuovo_cliente_input = st.text_input("Inserisci Nuova Ragione Sociale Cliente *", key="new_client_name_input")
    cliente = nuovo_cliente_input.strip() if nuovo_cliente_input else ""
else:
    cliente = cliente_selezionato_menu

# Campi di testo collegati direttamente al caricamento automatico
email_cliente = st.text_input("Email Cliente *", key="email_manual_input")
cellulare_cliente = st.text_input("Numero Cellulare Cliente *", key="cellulare_manual_input")

# === CAMPO INSERITO: PROPRIETARIO NUMERO / FIRMATARIO ===
firmatario_cliente = st.text_input("Nome di chi firma l'SMS (es. Sig. Mario Rossi) *")

marchio = st.text_input("Marchio Apparecchio *")
matricola = st.text_input("Matricola Apparecchio")
guasto_segnalato = st.text_area("Guasto Segnalato")
descrizione_lavori = st.text_area("Intervento Eseguito e Materiali Utilizzati *")
note_extra = st.text_area("Note Extra / Ricambi da ordinare")

km = st.number_input("Kilometri percorsi (Km)", min_value=0, value=0)
ore_lavoro = st.number_input("Ore di lavoro impiegate", min_value=0.0, value=0.0)
preventivo = st.radio("Richiedi Preventivo?", ["NO", "SI"])
urgente = st.radio("Intervento Urgente?", ["NO", "SI"])

# === CATTURA AUTOMATICA GEOLOCALIZZAZIONE GPS ===
st.write("📍 **Verifica Posizione GPS Intervento**")
loc = get_geolocation()
link_maps_str = "Posizione GPS Non Disponibile"
if loc and 'coords' in loc:
    lat = loc['coords']['latitude']
    lon = loc['coords']['longitude']
    link_maps_str = f"https://google.com{lat},{lon}"
    st.success(f"✅ Coordinate GPS Acquisite Correttamente!")
else:
    st.info("ℹ️ Consenti l'accesso alla geolocalizzazione se richiesto dal telefono per tracciare la firma d'intervento.")

file_immagini_caricate = st.file_uploader("📸 Carica o Scatta Foto dell'Intervento (Massimo 4 foto)", type=["png", "jpg", "jpeg"], accept_multiple_files=True)

# --- 2. GESTIONE SMS ---
st.subheader("🔒 Firma Digitale SMS Cliente")
if st.button("📲 INVIA CODICE DI VALIDAZIONE VIA SMS"):
    if not tecnico_autorizzato:
        st.error("⛔ AZIONE BLOCCATA: Devi prima selezionare il tuo nome Tecnico e inserire il PIN corretto in alto!")
    elif cliente_selezionato_menu == "➕ AGGIUNGI NUOVO CLIENTE" and not cliente:
        st.error("⚠️ Scrivi il nome del nuovo cliente nella casella di testo prima di inviare l'SMS!")
    elif not firmatario_cliente:
        st.error("⚠️ Inserisci il nome della persona fisica che firmerà l'SMS!")
    elif not cellulare_cliente or not cliente:
        st.error("⚠️ Inserisci Cliente e Cellulare!")
    else:
        st.session_state["codice_sms"] = str(random.randint(1000, 9999))
        st.session_state["sms_validato"] = False
        st.session_state["mostra_download"] = False
        
        from twilio.rest import Client
        
        ACCOUNT_SID = st.secrets["TWILIO_ACCOUNT_SID"]
        AUTH_TOKEN = st.secrets["TWILIO_AUTH_TOKEN"]
        NUMERO_TWILIO = st.secrets["TWILIO_NUMBER"]
        
        num_destinatario = cellulare_cliente.strip().replace(" ", "")
        if not num_destinatario.startswith("+"):
            if num_destinatario.startswith("39"):
                num_destinatario = "+" + num_destinatario
            else:
                num_destinatario = "+39" + num_destinatario
                
        testo_messaggio = f"Nova Servimpianti: Il tuo codice segreto di firma per l'intervento odierno e': {st.session_state['codice_sms']}"
        
        try:
            twilio_client = Client(ACCOUNT_SID, AUTH_TOKEN)
            twilio_client.messages.create(body=testo_messaggio, from_=NUMERO_TWILIO, to=num_destinatario)
            st.success(f"📩 SMS inviato correttamente al numero {num_destinatario}!")
        except Exception as e:
            st.error(f"❌ Errore critico Twilio: {e}")

if st.session_state["codice_sms"] is not None:
    st.info(f"👉 CODICE DI VALIDAZIONE D'EMERGENZA: {st.session_state['codice_sms']}")
    if not st.session_state["sms_validato"]:
        codice_inserito = st.text_input("Inserisci le 4 cifre:", key="codice_verifica_sms")
        if st.button("✅ VALIDA CODICE SMS"):
            if codice_inserito == st.session_state["codice_sms"]:
                st.session_state["sms_validato"] = True
                st.success("🔒 Validato!")
                st.rerun()
            else:
                st.error("❌ Codice errato!")
    else:
        st.success("🔒 Convalidato con Successo!")

# --- 3. LOGICA INVIO COPIA EMAIL GRAFICA IN HTML ---
def invia_email_pdf(destinatario, allegato_path, nome_cliente):
    email_mittente = "franciosoandrea@gmail.com" 
    password_mittente = st.secrets["GMAIL_PASSWORD"]
    
    msg_cli = MIMEMultipart('alternative')
    msg_cli['From'] = email_mittente
    msg_cli['To'] = destinatario
    msg_cli['Subject'] = f"Rapporto Intervento Ufficiale - Nova Servimpianti"
    
    html_cliente = f"""
    <html>
    <body style="font-family: 'Segoe UI', Arial, sans-serif; color: #333333; margin: 0; padding: 0; background-color: #F7FAFC;">
        <table align="center" border="0" cellpadding="0" cellspacing="0" width="100%" style="max-width: 600px; background-color: #ffffff; border: 1px solid #E2E8F0; margin-top: 20px; border-radius: 8px; overflow: hidden; box-shadow: 0 4px 6px rgba(0,0,0,0.05);">
            <tr>
                <td bgcolor="#1A365D" style="padding: 25px; text-align: center;">
                    <h1 style="color: #ffffff; margin: 0; font-size: 22px; letter-spacing: 1px;">NOVA SERVIMPIANTI SRLS</h1>
                    <p style="color: #90CDF4; margin: 5px 0 0 0; font-size: 13px;">Rapporto di Intervento Técnico Ufficiale</p>
                </td>
            </tr>
            <tr>
                <td style="padding: 30px;">
                    <p style="font-size: 16px; line-height: 24px; margin-top: 0;">Gentile Cliente,</p>
                    <p style="font-size: 15px; line-height: 24px;">Con la presente Le inviamo in allegato il <b>Rapporto d'Intervento Tecnico ufficiale</b> in formato PDF, relativo ai lavori eseguiti presso la Sua sede per l'azienda <b>{nome_cliente}</b>.</p>
                    <div style="background-color: #EDF2F7; border-left: 4px solid #2C5282; padding: 15px; margin: 25px 0; border-radius: 4px;">
                        <h3 style="margin: 0 0 10px 0; color: #2C5282; font-size: 14px; text-transform: uppercase; letter-spacing: 0.5px;">Riepilogo Documento</h3>
                        <table width="100%" style="font-size: 14px; border-collapse: collapse;">
                            <tr>
                                <td style="padding: 5px 0; color: #4A5568;" width="40%"><b>Stato Intervento:</b></td>
                                <td style="padding: 5px 0; color: #1A202C;">■ Chiuso e Convalidato</td>
                            </tr>
                            <tr>
                                <td style="padding: 5px 0; color: #4A5568;"><b>Firma Digitale:</b></td>
                                <td style="padding: 5px 0; color: #1A202C;">✓ Verificata via SMS OTP</td>
                            </tr>
                        </table>
                    </div>
                    <p style="font-size: 14px; line-height: 22px; color: #718096;">Troverà tutti i dettagli analitici direttamente all'interno del <b>file PDF allegato</b> a questa email. Per qualsiasi chiarimento non esitate a contattarci porgiamo distinti saluti .</p>
                </td>
            </tr>
            <tr>
                <td bgcolor="#F7FAFC" style="padding: 20px; text-align: center; border-top: 1px solid #E2E8F0; font-size: 12px; color: #718096;">
                    <b>Nova Servimpianti Srls</b><br/>Email Tecnica: franciosoandrea@me.com
                </td>
            </tr>
        </table>
    </body>
    </html>
    """
    msg_cli.attach(MIMEText(html_cliente, 'html'))
    
    msg_teco = MIMEMultipart()
    msg_teco['From'] = email_mittente
    msg_teco['To'] = "franciosoandrea@me.com"
    msg_teco['Subject'] = f"Nova Servimpianti - Backup Intervento: {nome_cliente}"
    msg_teco.attach(MIMEText("Rapporto registrato correttamente nel database.\nIn allegato trovi il PDF dell'intervento e il file Excel Generale aggiornato.", 'plain'))
    
    try:
        with open(allegato_path, "rb") as att_pdf:
            part_pdf = MIMEBase("application", "octet-stream")
            part_pdf.set_payload(att_pdf.read())
            encoders.encode_base64(part_pdf)
            part_pdf.add_header("Content-Disposition", f"attachment; filename= {allegato_path}")
            msg_cli.attach(part_pdf)
            
            part_pdf_teco = MIMEBase("application", "octet-stream")
            att_pdf.seek(0)
            part_pdf_teco.set_payload(att_pdf.read())
            encoders.encode_base64(part_pdf_teco)
            part_pdf_teco.add_header("Content-Disposition", f"attachment; filename= {allegato_path}")
            msg_teco.attach(part_pdf_teco)
            
        if os.path.exists(EXCEL_FILE):
            with open(EXCEL_FILE, "rb") as att_ex:
                part_ex = MIMEBase("application", "octet-stream")
                part_ex.set_payload(att_ex.read())
                encoders.encode_base64(part_ex)
                part_ex.add_header("Content-Disposition", f"attachment; filename= {EXCEL_FILE}")
                msg_teco.attach(part_ex)
                
        server = smtplib.SMTP("smtp.gmail.com", 587)
        server.starttls()
        server.login(email_mittente, password_mittente)
        server.sendmail(email_mittente, destinatario, msg_cli.as_string())
        server.sendmail(email_mittente, "franciosoandrea@me.com", msg_teco.as_string())
        server.quit()
        st.success("✉️ Documenti inviati! Email grafica con PDF inviata al cliente, PDF + Excel Storico inviati a franciosoandrea@me.com")
    except Exception as e:
        st.warning(f"⚠️ Nota: File registrati, ma l'invio email ha riscontrato un problema: {e}")
     
# --- 4. CREAZIONE PDF CON LINK GOOGLE MAPS E FOTO MULTIPLE ---
def elabora_pdf(pdf_filename, data_str, cliente, email_cliente, cellulare_cliente, firmatario, marchio, matricola, km, ore_lavoro, preventivo, urgent, guasto_segnalato, descrizione_lavori, note_extra, lista_file_immagini, stringa_firma, firma_tecnico, link_maps):
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, Image as RLImage
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import letter
    
    doc = SimpleDocTemplate(pdf_filename, pagesize=letter, leftMargin=40, rightMargin=40, topMargin=40, bottomMargin=40)
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle('T1', parent=styles['Heading1'], fontSize=18, textColor=colors.HexColor("#1A365D"), alignment=1, spaceAfter=20)
    section_heading = ParagraphStyle('T2', parent=styles['Heading3'], fontSize=12, textColor=colors.HexColor("#2C5282"), spaceBefore=14, spaceAfter=6)
    body_style = ParagraphStyle('T3', parent=styles['Normal'], fontSize=10, leading=16)
    firma_style = ParagraphStyle('T4', parent=styles['Normal'], fontSize=9, leading=14, textColor=colors.HexColor("#4A5568"))
    
    story = []
    if os.path.exists(LOGO_FILE):
        story.append(RLImage(LOGO_FILE, width=545, height=90))
        story.append(Spacer(1, 15))
        
    story.append(Paragraph("<b>RAPPORTO DI INTERVENTO TECNICO</b>", title_style))
    
    # === STRUTTURA DESIGN CON TITOLI IN BLU (#1A365D) E VALORI IN NERO ===
    dati_tabella = [
        [Paragraph(f"<font color='#1A365D'><b>Data Intervento:</b></font> {data_str}", body_style), 
         Paragraph(f"<font color='#1A365D'><b>Marchio Apparecchio:</b></font> {marchio}", body_style)],
         
        [Paragraph(f"<font color='#1A365D'><b>Cliente / Ragione Sociale:</b></font> {cliente}", body_style), 
         Paragraph(f"<font color='#1A365D'><b>Matricola:</b></font> {matricola if matricola else 'N.D.'}", body_style)],
         
        [Paragraph(f"<font color='#1A365D'><b>Email Cliente:</b></font> {email_cliente}", body_style), 
         Paragraph(f"<font color='#1A365D'><b>Kilometri Percorsi:</b></font> {km} Km", body_style)],
         
        [Paragraph(f"<font color='#1A365D'><b>Numero Cellulare:</b></font> {cellulare_cliente}", body_style), 
         Paragraph(f"<font color='#1A365D'><b>Ore Lavoro Impiegate:</b></font> {ore_lavoro}", body_style)],
         
        [Paragraph(f"<font color='#1A365D'><b>Firmatario/Collaboratore:</b></font> {firmatario if firmatario else 'N.D.'}", body_style), 
         Paragraph(f"<font color='#1A365D'><b>Richiede Preventivo:</b></font> {preventivo} &nbsp;&nbsp;|&nbsp;&nbsp; <font color='#1A365D'><b>Urgente:</b></font> {urgent}", body_style)]
    ]
    
    # Ripartizione esatta della larghezza (530 pixel utili)
    tabella_dati = Table(dati_tabella, colWidths=[265, 265])
    
    # Stile tabella con spaziatura ariosa e linee grigie eleganti sotto ogni riga
    tabella_dati.setStyle(TableStyle([
        ('VALIGN', (0, 0), (-1, -1), 'TOP'),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 6),
        ('TOPPADDING', (0, 0), (-1, -1), 6),
        ('LINEBELOW', (0, 0), (-1, -1), 0.5, colors.HexColor("#E2E8F0")),
    ]))
    
    story.append(tabella_dati)
    story.append(Spacer(1, 15))
    
    story.append(Paragraph("<b>■ GUASTO SEGNALATO</b>", section_heading))
    story.append(Paragraph(guasto_segnalato if guasto_segnalato else "N.D.", body_style))
    
    story.append(Paragraph("<b>■ LAVORI ESEGUITI E MATERIALI UTILIZZATI</b>", section_heading))
    story.append(Paragraph(descrizione_lavori, body_style))
    
    # === SEZIONE NOTE STAMPATA NEL PDF ===
    if note_extra:
        story.append(Paragraph("<b>■ NOTE EXTRA / RACCOMANDAZIONI</b>", section_heading))
        story.append(Paragraph(note_extra, body_style))
        
    story.append(Spacer(1, 15))
    
    story.append(Paragraph("<b>Firma del Tecnico Responsabile:</b>", body_style))
    story.append(Paragraph(f"<i>■ Convalidato e Firmato dal Tecnico: {firma_tecnico} il {data_str}</i>", firma_style))
    
    if "https" in str(link_maps):
        story.append(Paragraph(f"📍 <u><a href='{link_maps}' color='#2C5282'>■ Clicca qui per verificare la posizione GPS del Tecnico su Google Maps</a></u>", firma_style))
    else:
        story.append(Paragraph(f"📍 <i>Posizione GPS: Non disponibile o non autorizzata</i>", firma_style))
        
    story.append(Spacer(1, 15))
    
    story.append(Paragraph("<b>Firma per Accettazione Cliente:</b>", body_style))
    story.append(Paragraph(f"<i>■ {stringa_firma}</i>", firma_style))
    
    if lista_file_immagini and len(lista_file_immagini) > 0:
        story.append(Spacer(1, 15))
        story.append(Paragraph("<b>■ DOCUMENTAZIONE FOTOGRAFICA APPARECCHIO</b>", section_heading))
        for idx, file_img in enumerate(lista_file_immagini[:4]):
            story.append(Spacer(1, 10))
            foto_img = Image.open(file_img)
            foto_img.thumbnail((500, 350))
            temp_path = f"temp_allegato_{idx}.png"
            foto_img.save(temp_path)
            story.append(RLImage(temp_path, width=440, height=280))
            
    doc.build(story)

# --- 5. FUNZIONE GENERALE DI SCRITTURA DATI CON LINK GOOGLE MAPS E NOTE SU EXCEL ---
def registra_dati_intervento(data_str, tecnico, cliente, email_cliente, cellulare_cliente, firmatario, marchio, matricola, guasto_segnalato, descrizione_lavori, note_extra, km, ore_lavoro, preventivo, urgente, stringa_firma, link_maps):
    riga = {
        "Data Intervento": data_str,
        "Tecnico Responsabile": tecnico,
        "Ragione Sociale Cliente": cliente,
        "Email Cliente": email_cliente,
        "Cellulare Cliente": cellulare_cliente,
        "Firmatario / Proprietario Numero": firmatario if firmatario else "N.D.",
        "Marchio Apparecchio": marchio,
        "Matricola": matricola if matricola else "N.D.",
        "Guasto Segnalato": guasto_segnalato if guasto_segnalato else "N.D.",
        "Intervento Eseguito e Materiali Utilizzati": descrizione_lavori,
        "Note Extra": note_extra if note_extra else "N.D.",
        "Km Percorsi": km,
        "Ore Lavoro": ore_lavoro,
        "Richiede Preventivo?": preventivo,
        "Intervento Urgente?": urgente,
        "Firma Cliente": stringa_firma,
        "Link Google Maps GPS": link_maps
    }
    
    nome_foglio = cliente.replace(" ", "_").replace("/", "_").replace("\\", "_").replace("?", "_").replace("*", "_")[:30]
    
    if os.path.exists(EXCEL_FILE):
        try:
            wb = load_workbook(EXCEL_FILE, read_only=True)
            fogli_presenti = wb.sheetnames
            wb.close()
            if nome_foglio in fogli_presenti:
                df_esistente = pd.read_excel(EXCEL_FILE, sheet_name=nome_foglio)
                df_nuovo = pd.concat([df_esistente, pd.DataFrame([riga])], ignore_index=True)
            else:
                df_nuovo = pd.DataFrame([riga])
        except Exception:
            df_nuovo = pd.DataFrame([riga])
    else:
        df_nuovo = pd.DataFrame([riga])
        
    # Ordine sistemato delle colonne inclusa la colonna corretta dei lavori eseguiti e del firmatario
    colonne_ordinate = [
        "Data Intervento", "Tecnico Responsabile", "Ragione Sociale Cliente", 
        "Email Cliente", "Cellulare Cliente", "Firmatario / Proprietario Numero", 
        "Marchio Apparecchio", "Matricola", "Guasto Segnalato", 
        "Intervento Eseguito e Materiali Utilizzati", "Note Extra", 
        "Km Percorsi", "Ore Lavoro", "Richiede Preventivo?", 
        "Intervento Urgente?", "Firma Cliente", "Link Google Maps GPS"
    ]
    df_nuovo = df_nuovo.reindex(columns=colonne_ordinate)
    
    modalita = 'a' if os.path.exists(EXCEL_FILE) else 'w'
    parametri_writer = {'engine': 'openpyxl', 'mode': modalita}
    if modalita == 'a':
        parametri_writer['if_sheet_exists'] = 'replace'
        
    with pd.ExcelWriter(EXCEL_FILE, **parametri_writer) as writer:
        df_nuovo.to_excel(writer, sheet_name=nome_foglio, index=False)
        worksheet = writer.sheets[nome_foglio]
        
        for col_idx, col in enumerate(worksheet.columns, start=1):
            max_len = 0
            col_letter = get_column_letter(col_idx)
            for cell in col:
                if cell.value is not None:
                    max_len = max(max_len, len(str(cell.value)))
            worksheet.column_dimensions[col_letter].width = max(max_len + 4, 12)


# --- 6. BOTTONE DI SALVATAGGIO FINALIZZATO ---
st.subheader("💾 Registrazione")
st.warning("⚠️ ATTENZIONE TECNICO: Almeno una foto della scheda o della macchina è OBBLIGATORIA per poter chiudere l'intervento!")

if st.button("💾 REGISTRA E GENERA REPORT COMPLETO"):
    gps_final_link = link_maps_str if 'link_maps_str' in globals() else "Posizione GPS Non Disponibile"
    
    if cliente_selezionato_menu == "➕ AGGIUNGI NUOVO CLIENTE" and not cliente:
        st.error("❌ BLOCCO: Scrivi il nome del nuovo cliente nella casella di testo prima di salvare!")
    elif not cliente or not marchio or not descrizione_lavori or not email_cliente or not firmatario_cliente:
        st.error("⚠️ Compila i campi obbligatori (*) inclusa la persona che firma l'SMS!")
    elif not file_immagini_caricate or len(file_immagini_caricate) == 0:
        st.error("❌ BLOCCO: Devi caricare o scattare almeno 1 foto prima di salvare il report!")
    elif len(file_immagini_caricate) > 4:
        st.error("❌ BLOCCO: Puoi caricare al massimo 4 foto per ogni intervento!")
    elif not tecnico_autorizzato:
        st.error("⚠️ Il Tecnico deve inserire un PIN valido in alto per procedere!")
    elif not st.session_state["sms_validato"]:
        st.error("⚠️ Valida prima il codice SMS del cliente!")
    else:
        data_str = data_corrente.strftime("%d/%m/%Y")
        firma_tecnico_str = tecnico_selezionato
        
        # Stringa firma aggiornata con il nome del collaboratore/proprietario inserito a schermo
        stringa_firma_cli = f"Firmato via SMS OTP da {firmatario_cliente} (Cell: {cellulare_cliente}) il {data_str} (ID-{st.session_state['codice_sms']})"
        
        registra_dati_intervento(data_str, tecnico_selezionato, cliente, email_cliente, cellulare_cliente, firmatario_cliente, marchio, matricola, guasto_segnalato, descrizione_lavori, note_extra, km, ore_lavoro, preventivo, urgente, stringa_firma_cli, gps_final_link)
        
        c_pulito = cliente.replace(" ", "_").replace("/", "_")
        pdf_filename = f"Report_{data_corrente.strftime('%Y%m%d')}_{c_pulito}.pdf"
        st.session_state["ultimo_pdf"] = pdf_filename
        
        elabora_pdf(pdf_filename, data_str, cliente, email_cliente, cellulare_cliente, firmatario_cliente, marchio, matricola, km, ore_lavoro, preventivo, urgente, guasto_segnalato, descrizione_lavori, note_extra, file_immagini_caricate, stringa_firma_cli, firma_tecnico_str, gps_final_link)
        
        st.success("🎉 Registrato correttamente!")
        st.session_state["mostra_download"] = True
        invia_email_pdf(email_cliente, pdf_filename, cliente)
        st.rerun()

# --- 7. DOWNLOAD FISSI IN CODA ---
if st.session_state["mostra_download"]:
    st.subheader("📥 Scarica i File")
    with open(EXCEL_FILE, "rb") as f_ex:
        st.download_button("📥 Scarica Registro Excel", f_ex, file_name=EXCEL_FILE, key="b_ex")
    if st.session_state["ultimo_pdf"] and os.path.exists(st.session_state["ultimo_pdf"]):
        with open(st.session_state["ultimo_pdf"], "rb") as f_pd:
            st.download_button("📥 Scarica Questo PDF", f_pd, file_name=st.session_state["ultimo_pdf"], key="b_pd")
